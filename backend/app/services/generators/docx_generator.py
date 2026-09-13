from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
)
from docx.shared import (
    Inches,
    Pt,
)

from app.services.generators.content_utils import (
    clean_text,
    ensure_parent_directory,
    get_content_json,
    humanize_key,
    output_title,
)


class DOCXGenerator:

    def _append_value(
        self,
        document: Document,
        value: Any,
        level: int = 2,
    ) -> None:
        if value is None:
            return

        if isinstance(
            value,
            dict,
        ):
            for key, nested in (
                value.items()
            ):
                heading_level = min(
                    max(
                        level,
                        1,
                    ),
                    9,
                )

                document.add_heading(
                    humanize_key(
                        str(key)
                    ),
                    level=heading_level,
                )

                self._append_value(
                    document,
                    nested,
                    level=min(
                        level + 1,
                        9,
                    ),
                )

            return

        if isinstance(
            value,
            list,
        ):
            for item in value:
                if isinstance(
                    item,
                    dict,
                ):
                    paragraph = (
                        document.add_paragraph()
                    )

                    paragraph.add_run(
                        "•"
                    ).bold = True

                    self._append_value(
                        document,
                        item,
                        level=min(
                            level + 1,
                            9,
                        ),
                    )

                elif isinstance(
                    item,
                    list,
                ):
                    self._append_value(
                        document,
                        item,
                        level,
                    )

                else:
                    text = clean_text(
                        item
                    )

                    if text:
                        document.add_paragraph(
                            text,
                            style="List Bullet",
                        )

            return

        text = clean_text(
            value
        )

        if text:
            document.add_paragraph(
                text
            )

    def generate(
        self,
        transformation: dict,
        outputs: list[dict],
        destination: Path,
    ) -> Path:
        ensure_parent_directory(
            destination
        )

        document = Document()

        section = document.sections[0]

        section.top_margin = Inches(
            0.65
        )

        section.bottom_margin = Inches(
            0.65
        )

        section.left_margin = Inches(
            0.75
        )

        section.right_margin = Inches(
            0.75
        )

        styles = document.styles

        styles["Normal"].font.name = (
            "Aptos"
        )

        styles["Normal"].font.size = Pt(
            10.5
        )

        title = (
            transformation.get(
                "title"
            )
            or (
                "GenAI Content "
                "Transformation"
            )
        )

        title_paragraph = (
            document.add_heading(
                title,
                level=0,
            )
        )

        title_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        metadata = (
            document.add_paragraph()
        )

        metadata.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        metadata.add_run(
            (
                f"Audience: "
                f"{transformation.get('target_audience') or 'General'}"
                f" | Tone: "
                f"{transformation.get('tone') or 'Professional'}"
                f" | Language: "
                f"{transformation.get('language') or 'English'}"
            )
        )

        analysis = (
            transformation.get(
                "analysis_json"
            )
            or {}
        )

        if analysis:
            document.add_heading(
                "Content Intelligence",
                level=1,
            )

            important_analysis = {
                key: analysis.get(key)
                for key in (
                    "one_line_summary",
                    "executive_context",
                    "key_topics",
                    "key_facts",
                    "metrics",
                    "risks",
                    "opportunities",
                )
                if analysis.get(key)
            }

            self._append_value(
                document,
                important_analysis,
                level=2,
            )

            document.add_page_break()

        for index, output in (
            enumerate(outputs)
        ):
            document.add_heading(
                output_title(
                    output
                ),
                level=1,
            )

            content = (
                get_content_json(
                    output
                )
            )

            self._append_value(
                document,
                content,
                level=2,
            )

            if (
                index
                < len(outputs) - 1
            ):
                document.add_page_break()

        core_properties = (
            document.core_properties
        )

        core_properties.title = title

        core_properties.subject = (
            "SIH26154 GenAI "
            "Content Transformation"
        )

        core_properties.author = (
            "SIH26154 Platform"
        )

        document.save(
            str(destination)
        )

        return destination


docx_generator = DOCXGenerator()