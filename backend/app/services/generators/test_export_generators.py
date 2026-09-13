import csv
import json
from pathlib import Path

from docx import Document
from pptx import Presentation

from app.schemas.export import (
    ExportFormat,
)
from app.services.generators.registry import (
    generate_export_file,
)


def sample_transformation() -> dict:
    return {
        "id": (
            "11111111-1111-1111-"
            "1111-111111111111"
        ),
        "title": (
            "SIH26154 Demo "
            "Transformation"
        ),
        "status": "completed",
        "target_audience": (
            "Government decision makers"
        ),
        "tone": "Professional",
        "language": "English",
        "detail_level": "balanced",
        "objective": (
            "Explain the source clearly."
        ),
        "analysis_json": {
            "one_line_summary": (
                "A sample transformation."
            ),
            "key_topics": [
                "Generative AI",
                "Content Transformation",
            ],
            "risks": [
                "Incorrect source interpretation"
            ],
            "opportunities": [
                "Automated multi-format content"
            ],
        },
    }


def sample_outputs() -> list[dict]:
    return [
        {
            "id": "output-1",
            "output_type": (
                "executive_summary"
            ),
            "title": (
                "Executive Summary"
            ),
            "content_json": {
                "title": (
                    "Executive Summary"
                ),
                "summary": (
                    "Generative AI can "
                    "transform source content "
                    "into multiple formats."
                ),
                "key_points": [
                    "Multi-format generation",
                    "Audience-aware output",
                    "Structured extraction",
                ],
                "key_takeaway": (
                    "One source can generate "
                    "multiple communication assets."
                ),
            },
        },
        {
            "id": "output-2",
            "output_type": (
                "presentation"
            ),
            "title": (
                "Presentation"
            ),
            "content_json": {
                "title": (
                    "GenAI Transformation"
                ),
                "subtitle": (
                    "SIH26154"
                ),
                "slides": [
                    {
                        "slide_number": 1,
                        "title": (
                            "Problem"
                        ),
                        "bullets": [
                            (
                                "Content exists "
                                "in many formats."
                            ),
                            (
                                "Manual transformation "
                                "takes time."
                            ),
                        ],
                        "speaker_notes": (
                            "Explain the core problem."
                        ),
                        "visual_suggestion": (
                            "Input-to-output diagram"
                        ),
                    },
                    {
                        "slide_number": 2,
                        "title": (
                            "Solution"
                        ),
                        "bullets": [
                            (
                                "AI analysis"
                            ),
                            (
                                "RAG grounding"
                            ),
                            (
                                "Multi-format output"
                            ),
                        ],
                        "speaker_notes": (
                            "Explain the platform."
                        ),
                        "visual_suggestion": (
                            "Architecture diagram"
                        ),
                    },
                ],
                "closing_message": (
                    "Transform once, "
                    "communicate everywhere."
                ),
            },
        },
        {
            "id": "output-3",
            "output_type": (
                "video_script"
            ),
            "title": (
                "Video Script"
            ),
            "content_json": {
                "title": (
                    "SIH26154 Demo"
                ),
                "estimated_duration_seconds": (
                    10
                ),
                "opening_hook": (
                    "What if one document "
                    "could become everything?"
                ),
                "scenes": [
                    {
                        "scene_number": 1,
                        "duration_seconds": 5,
                        "visual_description": (
                            "Document upload screen"
                        ),
                        "voiceover": (
                            "Upload your source content."
                        ),
                        "on_screen_text": (
                            "Upload"
                        ),
                    },
                    {
                        "scene_number": 2,
                        "duration_seconds": 5,
                        "visual_description": (
                            "Outputs appearing"
                        ),
                        "voiceover": (
                            "Generate content for "
                            "multiple channels."
                        ),
                        "on_screen_text": (
                            "Transform"
                        ),
                    },
                ],
                "closing_cta": (
                    "Transform your content."
                ),
            },
        },
        {
            "id": "output-4",
            "output_type": (
                "structured_data"
            ),
            "title": (
                "Structured Data"
            ),
            "content_json": {
                "fields": [
                    {
                        "key": "Problem",
                        "value": (
                            "Manual content conversion"
                        ),
                        "category": (
                            "Challenge"
                        ),
                        "source_context": (
                            "Demo source"
                        ),
                    }
                ]
            },
        },
    ]


def test_generate_pdf(
    tmp_path: Path,
) -> None:
    destination = (
        tmp_path
        / "report.pdf"
    )

    generate_export_file(
        ExportFormat.PDF,
        sample_transformation(),
        sample_outputs(),
        destination,
    )

    assert destination.exists()

    assert (
        destination.stat().st_size
        > 100
    )

    with destination.open(
        "rb"
    ) as file:
        assert (
            file.read(4)
            == b"%PDF"
        )


def test_generate_docx(
    tmp_path: Path,
) -> None:
    destination = (
        tmp_path
        / "report.docx"
    )

    generate_export_file(
        ExportFormat.DOCX,
        sample_transformation(),
        sample_outputs(),
        destination,
    )

    assert destination.exists()

    document = Document(
        str(destination)
    )

    assert (
        len(document.paragraphs)
        > 0
    )


def test_generate_pptx(
    tmp_path: Path,
) -> None:
    destination = (
        tmp_path
        / "slides.pptx"
    )

    generate_export_file(
        ExportFormat.PPTX,
        sample_transformation(),
        sample_outputs(),
        destination,
    )

    assert destination.exists()

    presentation = (
        Presentation(
            str(destination)
        )
    )

    assert (
        len(
            presentation.slides
        )
        >= 3
    )


def test_generate_json(
    tmp_path: Path,
) -> None:
    destination = (
        tmp_path
        / "data.json"
    )

    generate_export_file(
        ExportFormat.JSON,
        sample_transformation(),
        sample_outputs(),
        destination,
    )

    payload = json.loads(
        destination.read_text(
            encoding="utf-8"
        )
    )

    assert (
        payload[
            "transformation"
        ][
            "status"
        ]
        == "completed"
    )

    assert len(
        payload[
            "outputs"
        ]
    ) == 4


def test_generate_csv(
    tmp_path: Path,
) -> None:
    destination = (
        tmp_path
        / "data.csv"
    )

    generate_export_file(
        ExportFormat.CSV,
        sample_transformation(),
        sample_outputs(),
        destination,
    )

    with destination.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        rows = list(
            csv.DictReader(
                file
            )
        )

    assert len(rows) > 0

    assert (
        "output_type"
        in rows[0]
    )


def test_generate_srt(
    tmp_path: Path,
) -> None:
    destination = (
        tmp_path
        / "captions.srt"
    )

    generate_export_file(
        ExportFormat.SRT,
        sample_transformation(),
        sample_outputs(),
        destination,
    )

    content = (
        destination.read_text(
            encoding="utf-8"
        )
    )

    assert (
        "00:00:00,000"
        in content
    )

    assert (
        "Upload your source content."
        in content
    )