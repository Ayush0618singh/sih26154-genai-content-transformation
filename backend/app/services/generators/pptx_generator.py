from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import (
    PP_ALIGN,
)
from pptx.util import (
    Inches,
    Pt,
)

from app.services.generators.content_utils import (
    clean_text,
    ensure_parent_directory,
    get_content_json,
    humanize_key,
    output_title,
    shorten_text,
)


class PPTXGenerator:

    def _configure_presentation(
        self,
        presentation: Presentation,
    ) -> None:
        presentation.slide_width = (
            Inches(13.333)
        )

        presentation.slide_height = (
            Inches(7.5)
        )

    def _add_title_slide(
        self,
        presentation: Presentation,
        transformation: dict,
    ) -> None:
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[0]
        )

        title = (
            transformation.get(
                "title"
            )
            or (
                "GenAI Content "
                "Transformation"
            )
        )

        if slide.shapes.title:
            slide.shapes.title.text = title

        subtitle = (
            slide.placeholders[1]
            if len(
                slide.placeholders
            ) > 1
            else None
        )

        if subtitle is not None:
            subtitle.text = (
                f"Target Audience: "
                f"{transformation.get('target_audience') or 'General'}\n"
                f"Tone: "
                f"{transformation.get('tone') or 'Professional'}\n"
                f"Language: "
                f"{transformation.get('language') or 'English'}"
            )

    def _set_notes(
        self,
        slide,
        notes: str,
    ) -> None:
        if not notes.strip():
            return

        try:
            notes_slide = (
                slide.notes_slide
            )

            text_frame = (
                notes_slide
                .notes_text_frame
            )

            text_frame.text = notes

        except Exception:
            pass

    def _add_content_slide(
        self,
        presentation: Presentation,
        title: str,
        bullets: list[str],
        notes: str = "",
    ) -> None:
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        if slide.shapes.title:
            slide.shapes.title.text = (
                shorten_text(
                    title,
                    100,
                )
            )

        body = (
            slide.placeholders[1]
            if len(
                slide.placeholders
            ) > 1
            else None
        )

        if body is None:
            return

        text_frame = body.text_frame

        text_frame.clear()

        cleaned_bullets = [
            shorten_text(
                bullet,
                220,
            )
            for bullet in bullets
            if bullet.strip()
        ][:7]

        if not cleaned_bullets:
            cleaned_bullets = [
                "Content available in source."
            ]

        for index, bullet in (
            enumerate(
                cleaned_bullets
            )
        ):
            paragraph = (
                text_frame.paragraphs[0]
                if index == 0
                else text_frame.add_paragraph()
            )

            paragraph.text = bullet
            paragraph.level = 0

            paragraph.font.size = Pt(
                20
            )

        self._set_notes(
            slide,
            notes,
        )

    def _presentation_output(
        self,
        outputs: list[dict],
    ) -> dict | None:
        for output in outputs:
            if (
                output.get(
                    "output_type"
                )
                == "presentation"
            ):
                return (
                    get_content_json(
                        output
                    )
                )

        return None

    def _value_to_bullets(
        self,
        value: Any,
    ) -> list[str]:
        bullets: list[str] = []

        if value is None:
            return bullets

        if isinstance(
            value,
            dict,
        ):
            for key, nested in (
                value.items()
            ):
                if isinstance(
                    nested,
                    (
                        dict,
                        list,
                    ),
                ):
                    bullets.append(
                        humanize_key(
                            str(key)
                        )
                    )

                    nested_bullets = (
                        self._value_to_bullets(
                            nested
                        )
                    )

                    bullets.extend(
                        nested_bullets[:3]
                    )

                else:
                    text = clean_text(
                        nested
                    )

                    if text:
                        bullets.append(
                            (
                                f"{humanize_key(str(key))}: "
                                f"{text}"
                            )
                        )

            return bullets

        if isinstance(
            value,
            list,
        ):
            for item in value:
                if isinstance(
                    item,
                    (
                        dict,
                        list,
                    ),
                ):
                    bullets.extend(
                        self._value_to_bullets(
                            item
                        )
                    )

                else:
                    text = clean_text(
                        item
                    )

                    if text:
                        bullets.append(
                            text
                        )

            return bullets

        text = clean_text(
            value
        )

        if text:
            bullets.append(
                text
            )

        return bullets

    def _build_from_ai_presentation(
        self,
        presentation: Presentation,
        presentation_content: dict,
    ) -> None:
        slides = (
            presentation_content.get(
                "slides"
            )
            or []
        )

        for slide_data in slides:
            if not isinstance(
                slide_data,
                dict,
            ):
                continue

            title = str(
                slide_data.get(
                    "title",
                    "Slide",
                )
            )

            bullets = [
                str(item)
                for item in (
                    slide_data.get(
                        "bullets"
                    )
                    or []
                )
            ]

            visual = str(
                slide_data.get(
                    "visual_suggestion",
                    "",
                )
            ).strip()

            notes = str(
                slide_data.get(
                    "speaker_notes",
                    "",
                )
            ).strip()

            if visual:
                notes = (
                    f"{notes}\n\n"
                    f"Visual suggestion: "
                    f"{visual}"
                ).strip()

            self._add_content_slide(
                presentation,
                title,
                bullets,
                notes,
            )

    def _build_generic(
        self,
        presentation: Presentation,
        outputs: list[dict],
    ) -> None:
        for output in outputs:
            content = (
                get_content_json(
                    output
                )
            )

            bullets = (
                self._value_to_bullets(
                    content
                )
            )

            self._add_content_slide(
                presentation,
                output_title(
                    output
                ),
                bullets,
            )

    def generate(
        self,
        transformation: dict,
        outputs: list[dict],
        destination: Path,
    ) -> Path:
        ensure_parent_directory(
            destination
        )

        presentation = Presentation()

        self._configure_presentation(
            presentation
        )

        self._add_title_slide(
            presentation,
            transformation,
        )

        presentation_content = (
            self._presentation_output(
                outputs
            )
        )

        if (
            presentation_content
            and presentation_content.get(
                "slides"
            )
        ):
            self._build_from_ai_presentation(
                presentation,
                presentation_content,
            )

        else:
            self._build_generic(
                presentation,
                outputs,
            )

        presentation.save(
            str(destination)
        )

        return destination


pptx_generator = PPTXGenerator()