from pathlib import Path

from docx import Document

from app.schemas.document import (
    DocumentExtractionResult,
    ExtractedPage,
)


async def parse_docx_document(
    file_path: Path,
    filename: str,
    mime_type: str,
) -> DocumentExtractionResult:

    document = Document(
        str(file_path)
    )

    sections: list[str] = []

    paragraph_count = 0

    for paragraph in (
        document.paragraphs
    ):

        text = paragraph.text.strip()

        if text:

            sections.append(
                text
            )

            paragraph_count += 1

    table_count = len(
        document.tables
    )

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):

        table_lines: list[str] = [
            f"[Table {table_index}]"
        ]

        for row in table.rows:

            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            table_lines.append(
                " | ".join(values)
            )

        sections.append(
            "\n".join(
                table_lines
            )
        )

    text = "\n\n".join(
        sections
    ).strip()

    page = ExtractedPage(
        page_number=1,
        text=text,
        character_count=len(text),
        used_ocr=False,
    )

    return DocumentExtractionResult(
        filename=filename,
        document_type="docx",
        mime_type=mime_type,
        extraction_method="python-docx",
        page_count=1,
        character_count=len(text),
        text=text,
        pages=[
            page,
        ],
        metadata={
            "paragraph_count": (
                paragraph_count
            ),
            "table_count": (
                table_count
            ),
        },
    )